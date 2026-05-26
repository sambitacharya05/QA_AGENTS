import os
import csv
from typing import Dict, List, Any
import openpyxl
from parsers.base import BaseParser
from parsers.id_utils import generate_node_id

class ExcelParser(BaseParser):
    """Parses Excel (.xlsx) and CSV files into clean tabular markdown structures."""
    
    def parse(self, file_path: str) -> Dict[str, Any]:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"File not found: {file_path}")
            
        file_name = os.path.basename(file_path)
        entities = []
        relationships = []
        raw_text_parts = []
        
        if file_path.endswith('.csv'):
            self._parse_csv(file_path, file_name, entities, raw_text_parts)
        else:
            self._parse_xlsx(file_path, file_name, entities, raw_text_parts)
            
        return {
            "raw_text": "\n".join(raw_text_parts),
            "entities": entities,
            "relationships": relationships
        }
        
    def _parse_csv(self, file_path: str, file_name: str, entities: List[Dict[str, Any]], raw_text_parts: List[str]) -> None:
        """Parses a raw CSV file."""
        rows = []
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            reader = csv.reader(f)
            for row in reader:
                rows.append([cell.strip() for cell in row])
                
        if not rows:
            return
            
        # Convert to Markdown table
        headers = rows[0]
        markdown_lines = []
        markdown_lines.append("| " + " | ".join(headers) + " |")
        markdown_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
        for row in rows[1:]:
            # Ensure row matches header length
            row_cells = row + [""] * (len(headers) - len(row))
            markdown_lines.append("| " + " | ".join(row_cells[:len(headers)]) + " |")
            
        markdown_str = "\n".join(markdown_lines)
        raw_text_parts.append(f"# CSV File: {file_name}\n{markdown_str}\n")
        
        node_id = generate_node_id("business_rule", file_name.replace('.csv', '').replace('.CSV', ''))
        entities.append({
            "id": node_id,
            "type": "business_rule",
            "name": f"CSV Data Grid ({file_name})",
            "description": f"Tabular rule grid extracted from CSV file: {file_name}.",
            "metadata": {
                "source_file": file_path,
                "headers": headers,
                "table_markdown": markdown_str
            }
        })
        
    def _parse_xlsx(self, file_path: str, file_name: str, entities: List[Dict[str, Any]], raw_text_parts: List[str]) -> None:
        """Parses an Excel (.xlsx) file, reading all sheets."""
        wb = openpyxl.load_workbook(file_path, data_only=True)
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            rows = []
            for row in ws.iter_rows(values_only=True):
                # Clean row values to strings, replacing None with empty
                cleaned_row = [str(val).strip() if val is not None else "" for val in row]
                # Keep only rows that aren't entirely empty
                if any(cleaned_row):
                    rows.append(cleaned_row)
                    
            if not rows:
                continue
                
            # Detect dimensions and convert to Markdown table
            headers = rows[0]
            markdown_lines = []
            markdown_lines.append("| " + " | ".join(headers) + " |")
            markdown_lines.append("| " + " | ".join(["---"] * len(headers)) + " |")
            for row in rows[1:]:
                row_cells = row + [""] * (len(headers) - len(row))
                markdown_lines.append("| " + " | ".join(row_cells[:len(headers)]) + " |")
                
            markdown_str = "\n".join(markdown_lines)
            raw_text_parts.append(f"# Sheet: {sheet_name} ({file_name})\n{markdown_str}\n")
            
            node_id = generate_node_id("business_rule", sheet_name)
            entities.append({
                "id": node_id,
                "type": "business_rule",
                "name": f"Excel Sheet: {sheet_name} ({file_name})",
                "description": f"Tabular grid sheet '{sheet_name}' representing insurance calculations, values, or pricing structures.",
                "metadata": {
                    "source_file": file_path,
                    "sheet_name": sheet_name,
                    "headers": headers,
                    "table_markdown": markdown_str
                }
            })
